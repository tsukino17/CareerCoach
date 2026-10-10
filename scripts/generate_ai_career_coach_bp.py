from copy import deepcopy
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN, MSO_VERTICAL_ANCHOR
from pptx.util import Inches, Pt


BASE_DIR = Path("/Users/sue/Documents/trae_projects/career/outputs/manual-20260513-115102/presentations/ai-career-coach-bp")
ASSET_DIR = BASE_DIR / "template-inspect" / "assets" / "ppt" / "media"
SOURCE_PPTX = Path("/Users/sue/Desktop/路演/产品核心介绍.pptx")
OUTPUT_PPTX = BASE_DIR / "output" / "AI职业教练-商业BP-可编辑版.pptx"

BLUE = RGBColor(78, 104, 245)
INK = RGBColor(30, 36, 49)
MUTED = RGBColor(96, 106, 128)
LIGHT = RGBColor(242, 245, 250)
BORDER = RGBColor(220, 225, 233)
PALE_BLUE = RGBColor(236, 241, 255)
PALE_ORANGE = RGBColor(255, 242, 232)
ORANGE = RGBColor(233, 126, 68)
WHITE = RGBColor(255, 255, 255)


def set_bg(slide, color=WHITE):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_textbox(slide, left, top, width, height, text, size=18, bold=False, color=INK,
                font_name="PingFang SC", align=PP_ALIGN.LEFT, linespacing=1.2):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = MSO_VERTICAL_ANCHOR.TOP
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = linespacing
    run = p.add_run()
    run.text = text
    font = run.font
    font.name = font_name
    font.size = Pt(size)
    font.bold = bold
    font.color.rgb = color
    return box


def add_kicker(slide, text, page=None):
    add_textbox(slide, Inches(0.75), Inches(0.42), Inches(1.0), Inches(0.22), "AI CAREER COACH",
                size=9, bold=True, color=MUTED)
    if page is not None:
        add_textbox(slide, Inches(12.05), Inches(0.42), Inches(0.35), Inches(0.22), f"{page:02d}",
                    size=9, color=MUTED, align=PP_ALIGN.RIGHT)


def add_title(slide, title, subtitle=None):
    add_textbox(slide, Inches(0.75), Inches(0.92), Inches(7.8), Inches(0.7), title, size=28, bold=True)
    if subtitle:
        add_textbox(slide, Inches(0.75), Inches(1.63), Inches(8.8), Inches(0.45), subtitle, size=11, color=MUTED)


def add_rule(slide, left, top, width, color=BORDER, height=0.01):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


def add_panel(slide, left, top, width, height, fill_color=LIGHT, line_color=None):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.adjustments[0] = 0.08
    if line_color:
        shape.line.color.rgb = line_color
        shape.line.width = Pt(1)
    else:
        shape.line.fill.background()
    return shape


def add_bullet_block(slide, left, top, width, heading, bullets, accent=BLUE):
    add_panel(slide, left, top, width, Inches(2.2), fill_color=WHITE, line_color=BORDER)
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, left, top, Inches(0.08), Inches(2.2))
    shape.fill.solid()
    shape.fill.fore_color.rgb = accent
    shape.line.fill.background()
    add_textbox(slide, left + Inches(0.18), top + Inches(0.16), width - Inches(0.32), Inches(0.35),
                heading, size=17, bold=True)
    y = top + Inches(0.62)
    for bullet in bullets:
        add_textbox(slide, left + Inches(0.22), y, width - Inches(0.36), Inches(0.42),
                    f"• {bullet}", size=11, color=MUTED)
        y += Inches(0.48)


def add_metric(slide, left, top, value, label, note, accent=BLUE):
    add_rule(slide, left, top, Inches(0.02), accent, Inches(0.68))
    add_textbox(slide, left + Inches(0.12), top - Inches(0.02), Inches(1.55), Inches(0.28), value,
                size=22, bold=True, color=INK)
    add_textbox(slide, left + Inches(0.12), top + Inches(0.3), Inches(1.55), Inches(0.2), label,
                size=9, bold=True, color=MUTED)
    add_textbox(slide, left + Inches(0.12), top + Inches(0.48), Inches(1.8), Inches(0.22), note,
                size=8, color=MUTED)


def add_image(slide, name, left, top, width=None, height=None):
    path = ASSET_DIR / name
    if width and height:
        slide.shapes.add_picture(str(path), left, top, width=width, height=height)
    elif width:
        slide.shapes.add_picture(str(path), left, top, width=width)
    elif height:
        slide.shapes.add_picture(str(path), left, top, height=height)
    else:
        slide.shapes.add_picture(str(path), left, top)


def add_footer(slide, text):
    add_rule(slide, Inches(0.75), Inches(6.7), Inches(11.1), BORDER)
    add_textbox(slide, Inches(0.75), Inches(6.78), Inches(10.8), Inches(0.18), text, size=8, color=MUTED)


def duplicate_slide(prs, index):
    source = prs.slides[index]
    blank = prs.slide_layouts[6]
    dest = prs.slides.add_slide(blank)
    for shape in source.shapes:
        el = deepcopy(shape.element)
        dest.shapes._spTree.insert_element_before(el, "p:extLst")
    return dest


def build_bp():
    OUTPUT_PPTX.parent.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # 1 cover
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    bar = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.78), Inches(1.8), Inches(0.08), Inches(1.0))
    bar.fill.solid()
    bar.fill.fore_color.rgb = BLUE
    bar.line.fill.background()
    add_textbox(slide, Inches(1.0), Inches(1.55), Inches(4.8), Inches(0.7), "AI职业教练", size=30, bold=True, color=BLUE)
    add_textbox(slide, Inches(1.0), Inches(2.26), Inches(7.2), Inches(0.7), "帮助用户在AI时代完成职业定位、能力迁移与行动转化", size=24, bold=True)
    add_textbox(slide, Inches(1.0), Inches(3.15), Inches(4.7), Inches(1.1),
                "一款结合教练式对话、职业分析与行动力工具箱的 AI 职业发展产品。",
                size=13, color=MUTED)
    add_metric(slide, Inches(1.0), Inches(5.22), "3大能力", "核心价值", "定位 / 分析 / 行动")
    add_metric(slide, Inches(3.1), Inches(5.22), "18页基础", "现有素材", "已完成产品逻辑表达")
    add_metric(slide, Inches(5.2), Inches(5.22), "MVP", "当前阶段", "核心功能已初步实现")
    add_image(slide, "image31.png", Inches(10.65), Inches(0.88), width=Inches(1.6))
    add_footer(slide, "基于原始《产品核心介绍》重组为商业 BP，可继续编辑。")

    # 2 problem
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_kicker(slide, "时代背景", 2)
    add_title(slide, "AI正在改写职业路径，传统职业辅导越来越难解决真实问题",
              "职业变化速度变快、岗位半衰期变短、用户更需要动态且可执行的陪伴式辅导。")
    add_bullet_block(slide, Inches(0.75), Inches(2.15), Inches(3.8), "岗位与技能快速失效", [
        "行业风口与岗位需求 2-5 年内频繁重构",
        "单一专业或固定岗位经验，不再天然等于长期竞争力",
        "用户需要识别可迁移、可复用的底层能力"
    ])
    add_bullet_block(slide, Inches(4.76), Inches(2.15), Inches(3.8), "职业规划普遍滞后", [
        "高校课程与真实岗位要求进一步脱节",
        "转岗、求职与重新进入市场的不适应感明显增强",
        "用户需要更个性化、更贴近现实场景的判断"
    ], accent=ORANGE)
    add_bullet_block(slide, Inches(8.77), Inches(2.15), Inches(3.8), "AI带来新的效率与焦虑", [
        "企业更关注轻量化、降本增效与结果导向",
        "个人一边看到机会，一边感到被替代的压力",
        "职业辅导不只要给方向，更要帮助用户开始行动"
    ])
    add_footer(slide, "原始材料来源：时代背景与痛点页。")

    # 3 market gap
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, LIGHT)
    add_kicker(slide, "用户痛点", 3)
    add_title(slide, "现有职业辅导的核心问题，不是信息不够，而是方法太慢、太贵、太模板化")
    add_bullet_block(slide, Inches(0.75), Inches(2.05), Inches(3.7), "流程繁琐", [
        "填表、测评、上传简历、等待分析，单次体验链路过长",
        "传统测评模型和职业推荐标签存在明显时代滞后"
    ])
    add_bullet_block(slide, Inches(4.82), Inches(2.05), Inches(3.7), "成本高、频次低", [
        "真人咨询通常 699-3000 元起，陪跑项目更高",
        "通常按周预约，无法承接用户高频、碎片化的决策需求"
    ], accent=ORANGE)
    add_bullet_block(slide, Inches(8.89), Inches(2.05), Inches(3.7), "推荐静态且同质化", [
        "容易围绕热门岗位给出建议，忽略用户个体差异",
        "缺少对用户长期发展方向与行动转换的持续支持"
    ])
    add_panel(slide, Inches(0.75), Inches(5.45), Inches(11.84), Inches(0.8), fill_color=WHITE, line_color=BORDER)
    add_textbox(slide, Inches(1.0), Inches(5.68), Inches(11.2), Inches(0.28),
                "这意味着市场缺的不是“更多职业信息”，而是一个能持续理解用户、动态推演方向并推动行动的 AI 职业教练。",
                size=14, bold=True)
    add_footer(slide, "原始材料来源：当前职业辅导产品为什么不能解决问题。")

    # 4 solution
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_kicker(slide, "解决方案", 4)
    add_title(slide, "AI职业教练把职业探索拆成三件事：看清自己、找到方向、推进现实行动")
    add_panel(slide, Inches(0.75), Inches(2.0), Inches(11.84), Inches(3.45), fill_color=PALE_BLUE)
    cols = [
        ("01", "看清自己", "通过共情式对话复现工作场景，挖掘用户的天赋、动机、优势与真实偏好。"),
        ("02", "找到方向", "结合能力迁移与时代趋势，输出更动态、更个性化的岗位与发展路径建议。"),
        ("03", "推进行动", "借助提示、日记、演练、模拟面试等工具，把认知转化为具体行动。"),
    ]
    x = Inches(1.1)
    for idx, (num, title, body) in enumerate(cols):
        add_textbox(slide, x, Inches(2.45), Inches(0.6), Inches(0.35), num, size=22, bold=True, color=BLUE)
        add_textbox(slide, x, Inches(2.92), Inches(2.8), Inches(0.35), title, size=18, bold=True)
        add_textbox(slide, x, Inches(3.38), Inches(2.85), Inches(1.2), body, size=11, color=MUTED)
        if idx < 2:
            add_rule(slide, x + Inches(3.15), Inches(2.55), Inches(0.55), BLUE, Inches(0.02))
        x += Inches(3.9)
    add_textbox(slide, Inches(0.75), Inches(5.86), Inches(11.5), Inches(0.3),
                "对用户而言，这不是一次性“测评结果”，而是一位可持续陪伴、不断校准和推动落地的 AI 教练。",
                size=12, color=MUTED)
    add_footer(slide, "结构重组自：产品理念、报告页、岗位推荐页与行动力工具箱页。")

    # 5 why now
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_kicker(slide, "产品理念", 5)
    add_title(slide, "产品判断是：在快速变化的时代，真正稳定的是能力、动机与天赋")
    add_panel(slide, Inches(0.75), Inches(2.0), Inches(3.68), Inches(3.2), fill_color=WHITE, line_color=BORDER)
    add_panel(slide, Inches(4.83), Inches(2.0), Inches(3.68), Inches(3.2), fill_color=WHITE, line_color=BORDER)
    add_panel(slide, Inches(8.91), Inches(2.0), Inches(3.68), Inches(3.2), fill_color=WHITE, line_color=BORDER)
    cards = [
        ("能力与特质", "与其只问“会不会某个工具”，更重要的是用户是否具备可迁移的审美、分析、表达、组织与落地能力。"),
        ("动机与热情", "职业选择背后有更深层的驱动力。看见用户真正想做什么，才能找到更宽的职业边界。"),
        ("天赋与特长", "长期优势往往来自做起来最顺手、最有正反馈的事情，而不只是热门岗位和短期薪资。"),
    ]
    x = Inches(1.0)
    for title, body in cards:
        add_textbox(slide, x, Inches(2.35), Inches(2.9), Inches(0.35), title, size=18, bold=True)
        add_textbox(slide, x, Inches(2.9), Inches(2.9), Inches(1.6), body, size=11, color=MUTED)
        x += Inches(4.08)
    add_panel(slide, Inches(0.75), Inches(5.55), Inches(11.84), Inches(0.82), fill_color=PALE_ORANGE)
    add_textbox(slide, Inches(1.0), Inches(5.8), Inches(11.2), Inches(0.26),
                "因此，AI职业教练的价值不只是“推荐岗位”，而是帮助用户识别自己的长期竞争力，并据此做出更好的职业决策。",
                size=13, bold=True)
    add_footer(slide, "原始材料来源：快速变化的AI时代什么是不变的。")

    # 6 methodology
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, LIGHT)
    add_kicker(slide, "理论与方法", 6)
    add_title(slide, "方法论不是装饰，而是产品可信度的一部分")
    methods = [
        ("福格行为模型", "行为 = 动机 + 能力 + 提示，用更低摩擦的方式推动用户开始行动。"),
        ("TRUTH转化模型", "围绕觉察、转化、行动，帮助用户把职业兴趣逐步落到现实。"),
        ("教练式对话", "不是直接告诉用户答案，而是通过启发式追问帮助其看见自身资源与渴望。"),
        ("焦点解决短程疗法", "帮助迷茫焦虑的用户重新聚焦已有优势，建立可行动的信心。"),
    ]
    positions = [
        (Inches(0.75), Inches(2.0)),
        (Inches(6.1), Inches(2.0)),
        (Inches(0.75), Inches(4.0)),
        (Inches(6.1), Inches(4.0)),
    ]
    for (title, body), (x, y) in zip(methods, positions):
        add_panel(slide, x, y, Inches(5.2), Inches(1.55), fill_color=WHITE, line_color=BORDER)
        add_textbox(slide, x + Inches(0.22), y + Inches(0.2), Inches(4.7), Inches(0.28), title, size=16, bold=True)
        add_textbox(slide, x + Inches(0.22), y + Inches(0.6), Inches(4.7), Inches(0.62), body, size=11, color=MUTED)
    add_footer(slide, "原始材料来源：心理学与教练对话指导。")

    # 7 product engine
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_kicker(slide, "核心能力", 7)
    add_title(slide, "产品底层由三层能力组成：对话理解、职业匹配、行动转化")
    add_panel(slide, Inches(0.75), Inches(2.05), Inches(11.84), Inches(3.95), fill_color=WHITE, line_color=BORDER)
    lanes = [
        ("自然对话层", "使用多模型调试教练式对话，强调自然、共情与个性化理解。", BLUE),
        ("匹配分析层", "根据用户天赋、经历、性格与发展偏好，推荐更适合的岗位、公司和行业。", ORANGE),
        ("行动工具层", "通过系统提示、行动日记、方案演练、模拟面试等工具帮助用户推进现实转化。", BLUE),
    ]
    y = Inches(2.42)
    for title, body, accent in lanes:
        band = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(1.0), y, Inches(2.0), Inches(0.82))
        band.fill.solid()
        band.fill.fore_color.rgb = accent
        band.line.fill.background()
        add_textbox(slide, Inches(1.2), y + Inches(0.22), Inches(1.6), Inches(0.24), title,
                    size=15, bold=True, color=WHITE)
        add_textbox(slide, Inches(3.25), y + Inches(0.14), Inches(8.6), Inches(0.52), body, size=12, color=MUTED)
        y += Inches(1.07)
    add_textbox(slide, Inches(1.0), Inches(5.48), Inches(10.9), Inches(0.28),
                "当前原始文档提到已使用 qwen-plus、gemini、deepseek 等模型进行训练与调试；这里保留为产品能力说明，不扩展未经验证的数据指标。",
                size=9.5, color=MUTED)
    add_footer(slide, "原始材料来源：AI智能底层驱动。")

    # 8 product experience - screenshot
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_kicker(slide, "产品演示", 8)
    add_title(slide, "用户体验从一段对话开始，而不是一张静态测评表")
    add_textbox(slide, Inches(0.75), Inches(1.68), Inches(4.2), Inches(0.8),
                "这一页保留真实产品界面素材，强化“这是已经可以被体验的产品”，而不是纯概念方案。",
                size=12, color=MUTED)
    add_panel(slide, Inches(0.75), Inches(2.3), Inches(4.2), Inches(2.85), fill_color=PALE_BLUE)
    add_textbox(slide, Inches(1.0), Inches(2.58), Inches(3.6), Inches(0.26), "体验特点", size=16, bold=True)
    points = [
        "以教练式提问帮助用户回到真实工作场景",
        "把“我适合什么工作”拆解为可追问、可验证的过程",
        "相比一次性测试，更适合高频迭代和持续校准",
    ]
    y = Inches(3.0)
    for p in points:
        add_textbox(slide, Inches(1.0), y, Inches(3.5), Inches(0.4), f"• {p}", size=11, color=MUTED)
        y += Inches(0.48)
    add_image(slide, "image28.png", Inches(6.35), Inches(1.95), height=Inches(4.95))
    add_footer(slide, "真实素材来源：原始PPT中的产品界面演示页。")

    # 9 output report
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, LIGHT)
    add_kicker(slide, "产品输出", 9)
    add_title(slide, "产品输出的价值，不是“给答案”，而是形成一份更可行动的职业分析")
    add_image(slide, "image30.png", Inches(0.9), Inches(2.0), height=Inches(4.75))
    add_panel(slide, Inches(7.2), Inches(2.12), Inches(5.0), Inches(4.1), fill_color=WHITE, line_color=BORDER)
    add_textbox(slide, Inches(7.48), Inches(2.38), Inches(4.3), Inches(0.28), "报告核心输出", size=16, bold=True)
    outputs = [
        "个人关键特质总结：提炼能力、目标、愿景与优势",
        "天赋适配岗位：结合趋势给出更长期的方向建议",
        "可复用能力：识别可迁移、可跨岗位使用的竞争力",
        "岗位对比与推荐：帮助用户做更清晰的选择",
    ]
    y = Inches(2.84)
    for item in outputs:
        add_textbox(slide, Inches(7.48), y, Inches(4.15), Inches(0.48), f"• {item}", size=11, color=MUTED)
        y += Inches(0.68)
    add_panel(slide, Inches(7.48), Inches(5.28), Inches(3.8), Inches(0.62), fill_color=PALE_ORANGE)
    add_textbox(slide, Inches(7.7), Inches(5.5), Inches(3.35), Inches(0.2),
                "重点是让用户理解“为什么适合”，而不是只看到岗位名称。", size=10.5, bold=True)
    add_footer(slide, "真实素材来源：深度职业分析报告、岗位推荐与对比页。")

    # 10 action loop
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_kicker(slide, "行动转化", 10)
    add_title(slide, "如果不能推动行动，职业辅导的价值会快速衰减")
    add_panel(slide, Inches(0.9), Inches(2.3), Inches(11.5), Inches(2.55), fill_color=WHITE, line_color=BORDER)
    stages = [
        ("觉察", "复盘经历，识别优势与卡点"),
        ("判断", "理解岗位本质、匹配度与方向"),
        ("演练", "模拟面试、方案演练、行动拆解"),
        ("推进", "通过提示与记录机制形成持续行动"),
    ]
    x = Inches(1.25)
    for idx, (title, body) in enumerate(stages):
        add_textbox(slide, x, Inches(2.68), Inches(1.6), Inches(0.3), title, size=17, bold=True)
        add_textbox(slide, x, Inches(3.08), Inches(1.75), Inches(0.75), body, size=10.5, color=MUTED)
        if idx < 3:
            add_rule(slide, x + Inches(1.95), Inches(3.02), Inches(0.56), BLUE, Inches(0.03))
        x += Inches(2.65)
    add_panel(slide, Inches(0.9), Inches(5.28), Inches(11.5), Inches(0.78), fill_color=PALE_BLUE)
    add_textbox(slide, Inches(1.15), Inches(5.52), Inches(10.9), Inches(0.22),
                "这也是产品后续研发重点：把职业认知结果，进一步转化为用户能执行、愿意执行、能持续执行的行动方案。",
                size=12, bold=True)
    add_footer(slide, "原始材料来源：行动力工具箱与项目进度说明。")

    # 11 roadmap / progress
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_kicker(slide, "项目进展", 11)
    add_title(slide, "目前已完成 MVP 核心链路，下一步聚焦产品深化与商业化验证")
    add_panel(slide, Inches(0.75), Inches(2.05), Inches(3.7), Inches(3.45), fill_color=WHITE, line_color=BORDER)
    add_textbox(slide, Inches(1.0), Inches(2.34), Inches(3.1), Inches(0.28), "已完成", size=16, bold=True)
    done = [
        "核心教练对话功能已初步实现",
        "职业分析报告与岗位推荐功能已打通",
        "产品已可通过网页链接进行体验",
    ]
    y = Inches(2.78)
    for item in done:
        add_textbox(slide, Inches(1.0), y, Inches(3.0), Inches(0.42), f"• {item}", size=11, color=MUTED)
        y += Inches(0.54)
    add_panel(slide, Inches(4.82), Inches(2.05), Inches(3.7), Inches(3.45), fill_color=WHITE, line_color=BORDER)
    add_textbox(slide, Inches(5.07), Inches(2.34), Inches(3.1), Inches(0.28), "下一阶段重点", size=16, bold=True)
    nexts = [
        "加强行动力转化与持续陪跑能力",
        "优化推荐机制与个性化判断质量",
        "验证适合的商业模式与增长路径",
    ]
    y = Inches(2.78)
    for item in nexts:
        add_textbox(slide, Inches(5.07), y, Inches(3.0), Inches(0.42), f"• {item}", size=11, color=MUTED)
        y += Inches(0.54)
    add_panel(slide, Inches(8.89), Inches(2.05), Inches(3.7), Inches(3.45), fill_color=PALE_ORANGE)
    add_textbox(slide, Inches(9.14), Inches(2.34), Inches(3.1), Inches(0.28), "当前诉求", size=16, bold=True)
    asks = [
        "AI模型训练与推荐算法指导",
        "心理学与职业教练方法论探讨",
        "商业化落地、品牌建设与用户增长支持",
    ]
    y = Inches(2.78)
    for item in asks:
        add_textbox(slide, Inches(9.14), y, Inches(3.0), Inches(0.42), f"• {item}", size=11, color=MUTED)
        y += Inches(0.54)
    add_footer(slide, "原始材料来源：项目试用链接、项目进度与所需资源页。")

    # 12 founder
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, LIGHT)
    add_kicker(slide, "团队与试用", 12)
    add_title(slide, "项目由创始人独立推进，已具备明确的问题意识、产品原型与方法论基础")
    add_panel(slide, Inches(0.75), Inches(2.02), Inches(7.4), Inches(3.95), fill_color=WHITE, line_color=BORDER)
    add_textbox(slide, Inches(1.0), Inches(2.35), Inches(2.0), Inches(0.3), "创始人：王旭", size=18, bold=True)
    founder = [
        "目前主要由创始人本人负责产品设计、模型训练和代码研发",
        "中科院心理所心理学与人工智能 2018 级在职研究生",
        "8 年在线 IT 职业教育课程设计与教研经历",
    ]
    y = Inches(2.88)
    for item in founder:
        add_textbox(slide, Inches(1.0), y, Inches(6.6), Inches(0.5), f"• {item}", size=11.5, color=MUTED)
        y += Inches(0.68)
    add_panel(slide, Inches(8.55), Inches(2.02), Inches(4.04), Inches(3.95), fill_color=WHITE, line_color=BORDER)
    add_textbox(slide, Inches(8.84), Inches(2.35), Inches(3.45), Inches(0.28), "试用信息", size=17, bold=True)
    add_textbox(slide, Inches(8.84), Inches(2.82), Inches(3.25), Inches(0.7),
                "产品链接：\nhttps://echotalent.fun/", size=11.5, color=MUTED)
    add_textbox(slide, Inches(8.84), Inches(3.72), Inches(3.25), Inches(0.7),
                "当前说明：\n手机和电脑端均可使用，海外服务器可能偶有不稳定。", size=11.5, color=MUTED)
    add_textbox(slide, Inches(8.84), Inches(4.72), Inches(3.25), Inches(0.8),
                "建议用途：\n用于体验核心教练对话、职业分析与岗位对比链路。", size=11.5, color=MUTED)
    add_footer(slide, "这版为可编辑商业 BP，可继续补融资、市场规模、商业模式与财务预测。")

    prs.save(OUTPUT_PPTX)


if __name__ == "__main__":
    build_bp()
