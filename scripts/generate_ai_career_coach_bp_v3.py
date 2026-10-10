from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import MSO_VERTICAL_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


BASE_DIR = Path("/Users/sue/Documents/trae_projects/career/outputs/manual-20260513-115102/presentations/ai-career-coach-bp")
ASSET_DIR = BASE_DIR / "template-inspect" / "assets" / "ppt" / "media"
OUTPUT_PPTX = BASE_DIR / "output" / "AI职业教练-商业BP-v3-精简融资版.pptx"

BG = RGBColor(250, 248, 244)
WHITE = RGBColor(255, 255, 255)
INK = RGBColor(30, 36, 48)
MUTED = RGBColor(110, 117, 132)
LINE = RGBColor(229, 229, 224)
BLUE = RGBColor(80, 118, 239)
BLUE_SOFT = RGBColor(235, 241, 255)
MINT = RGBColor(113, 190, 170)
MINT_SOFT = RGBColor(235, 247, 243)
PEACH = RGBColor(236, 161, 130)
PEACH_SOFT = RGBColor(252, 240, 234)
LAV = RGBColor(138, 135, 239)
LAV_SOFT = RGBColor(239, 238, 255)
GOLD = RGBColor(212, 160, 79)
GOLD_SOFT = RGBColor(248, 242, 226)


def bg(slide, color=BG):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def shape(slide, left, top, width, height, fill, line=None, rounded=True):
    kind = MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if rounded else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    shp = slide.shapes.add_shape(kind, left, top, width, height)
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if rounded:
        shp.adjustments[0] = 0.08
    if line:
        shp.line.color.rgb = line
        shp.line.width = Pt(1)
    else:
        shp.line.fill.background()
    return shp


def text(slide, left, top, width, height, value, *, size=18, bold=False, color=INK,
         align=PP_ALIGN.LEFT, font="PingFang SC", spacing=1.2):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = MSO_VERTICAL_ANCHOR.TOP
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = spacing
    run = p.add_run()
    run.text = value
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def header(slide, num, title, subtitle=None, width=9.6):
    text(slide, Inches(0.78), Inches(0.4), Inches(1.2), Inches(0.18), "AI CAREER COACH", size=9, bold=True, color=MUTED)
    text(slide, Inches(12.0), Inches(0.4), Inches(0.38), Inches(0.18), f"{num:02d}", size=9, color=MUTED, align=PP_ALIGN.RIGHT)
    text(slide, Inches(0.78), Inches(0.84), Inches(width), Inches(0.62), title, size=28, bold=True)
    if subtitle:
        text(slide, Inches(0.78), Inches(1.55), Inches(10.5), Inches(0.34), subtitle, size=11, color=MUTED)


def footer(slide, value):
    shape(slide, Inches(0.78), Inches(6.72), Inches(11.1), Inches(0.012), LINE, rounded=False)
    text(slide, Inches(0.78), Inches(6.8), Inches(11.0), Inches(0.14), value, size=8, color=MUTED)


def chip(slide, left, top, label, *, fill=WHITE, line=LINE, color=MUTED, size=8.5):
    width = max(0.72, 0.12 * len(label) + 0.38)
    shape(slide, left, top, Inches(width), Inches(0.28), fill, line=line, rounded=True)
    text(slide, left + Inches(0.12), top + Inches(0.05), Inches(width - 0.24), Inches(0.14), label, size=size, bold=True, color=color, align=PP_ALIGN.CENTER)


def panel(slide, left, top, width, height, *, fill=WHITE, line=LINE, title=None, title_color=INK):
    shape(slide, left, top, width, height, fill, line=line, rounded=True)
    if title:
        text(slide, left + Inches(0.2), top + Inches(0.18), width - Inches(0.4), Inches(0.22), title, size=16, bold=True, color=title_color)


def bullets(slide, left, top, width, items, *, size=11, color=MUTED, gap=0.46):
    y = top
    for item in items:
        text(slide, left, y, width, Inches(0.34), f"• {item}", size=size, color=color)
        y += Inches(gap)


def metric(slide, left, top, value, label, note, accent=BLUE):
    shape(slide, left, top, Inches(0.03), Inches(0.68), accent, rounded=False)
    text(slide, left + Inches(0.14), top - Inches(0.02), Inches(1.6), Inches(0.28), value, size=22, bold=True)
    text(slide, left + Inches(0.14), top + Inches(0.3), Inches(1.7), Inches(0.16), label, size=9, bold=True, color=MUTED)
    text(slide, left + Inches(0.14), top + Inches(0.48), Inches(1.85), Inches(0.14), note, size=8, color=MUTED)


def picture(slide, filename, left, top, *, width=None, height=None):
    path = ASSET_DIR / filename
    if width and height:
        slide.shapes.add_picture(str(path), left, top, width=width, height=height)
    elif width:
        slide.shapes.add_picture(str(path), left, top, width=width)
    elif height:
        slide.shapes.add_picture(str(path), left, top, height=height)
    else:
        slide.shapes.add_picture(str(path), left, top)


def line(slide, x1, y1, x2, y2, color=LINE, width=1.5):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    c.line.color.rgb = color
    c.line.width = Pt(width)
    return c


def arrow_stage(slide, left, top, width, fill, en, cn):
    shp = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.CHEVRON, left, top, width, Inches(0.96))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.fill.background()
    text(slide, left + Inches(0.16), top + Inches(0.16), width - Inches(0.32), Inches(0.16), en, size=9.5, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    text(slide, left + Inches(0.16), top + Inches(0.42), width - Inches(0.32), Inches(0.22), cn, size=17, bold=True, color=WHITE, align=PP_ALIGN.CENTER)


def build():
    OUTPUT_PPTX.parent.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # 1 cover
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    shape(s, Inches(8.02), Inches(0.52), Inches(4.6), Inches(6.2), WHITE, line=LINE)
    shape(s, Inches(0.78), Inches(5.92), Inches(11.0), Inches(0.42), PEACH_SOFT, line=PEACH_SOFT)
    text(s, Inches(0.78), Inches(0.44), Inches(1.4), Inches(0.18), "AI CAREER COACH", size=9, bold=True, color=MUTED)
    text(s, Inches(0.78), Inches(1.18), Inches(5.5), Inches(0.7), "AI职业教练", size=34, bold=True)
    text(s, Inches(0.78), Inches(1.98), Inches(6.1), Inches(0.86), "帮助用户在AI时代完成职业定位、能力迁移与行动转化", size=24, bold=True)
    text(s, Inches(0.78), Inches(3.02), Inches(4.8), Inches(0.84), "一款结合教练式对话、职业镜像报告、现实路径地图与行动推进机制的 AI 职业发展产品。", size=13, color=MUTED)
    metric(s, Inches(0.78), Inches(4.74), "MVP", "当前阶段", "核心链路已初步打通")
    metric(s, Inches(3.0), Inches(4.74), "现实判断", "核心差异", "不止给方向，更给路径")
    metric(s, Inches(5.22), Inches(4.74), "融资BP", "本次版本", "补齐路演与资金用途")
    picture(s, "image26.png", Inches(8.34), Inches(0.86), width=Inches(3.96))
    chip(s, Inches(8.46), Inches(5.42), "真实产品报告界面", fill=PEACH_SOFT, line=PEACH_SOFT, color=PEACH)
    text(s, Inches(8.46), Inches(5.78), Inches(3.5), Inches(0.36), "已经能够输出职业镜像报告，并逐步把认知翻译成现实职业判断。", size=10.5, color=MUTED)
    footer(s, "v3：精简融资版，强化视觉统一、留白与投资人阅读节奏。")

    # 2 problem
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    header(s, 2, "AI正在改写职业决策逻辑，用户需要的不只是信息，而是更快形成现实判断")
    panel(s, Inches(0.78), Inches(2.06), Inches(3.68), Inches(3.92), title="变化更快")
    bullets(s, Inches(1.02), Inches(2.66), Inches(3.0), [
        "岗位半衰期缩短，热点赛道与能力要求持续变化",
        "单一专业与固定履历不再天然构成长期竞争力",
        "用户需要识别可迁移、可复用的底层能力",
    ], gap=0.68)
    panel(s, Inches(4.82), Inches(2.06), Inches(3.68), Inches(3.92), title="滞后更明显", fill=BLUE_SOFT, line=BLUE_SOFT)
    bullets(s, Inches(5.06), Inches(2.66), Inches(3.0), [
        "高校课程与真实岗位要求进一步脱节",
        "转岗、求职、重返市场的人群更容易感到失配",
        "传统职业规划难以承接动态变化",
    ], gap=0.68)
    panel(s, Inches(8.86), Inches(2.06), Inches(3.68), Inches(3.92), title="焦虑更高")
    bullets(s, Inches(9.1), Inches(2.66), Inches(3.0), [
        "AI带来机会，同时放大被替代与落后的焦虑",
        "企业更强调轻量化、效率与结果导向",
        "职业辅导需要从“建议”升级为“陪伴式判断系统”",
    ], gap=0.68)
    footer(s, "原始材料来源：时代背景与用户痛点页。")

    # 3 gap
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    header(s, 3, "现有职业辅导产品的核心问题，不是没有内容，而是太慢、太贵、太模板化")
    panel(s, Inches(0.78), Inches(2.12), Inches(11.76), Inches(3.58))
    cols = [
        ("流程繁琐", "填表、测评、上传简历、等待分析，体验链路长，反馈时效慢。", GOLD_SOFT),
        ("成本高、频次低", "真人咨询价格不低，且多为按周预约，难承接高频职业决策。", PEACH_SOFT),
        ("推荐静态且同质化", "容易围绕热门岗位给出答案，忽略个体差异与长期潜力。", MINT_SOFT),
    ]
    for i, (h, b, fill) in enumerate(cols):
        x = Inches(1.08 + 3.8 * i)
        shape(s, x, Inches(2.42), Inches(3.28), Inches(2.74), fill, line=fill)
        text(s, x + Inches(0.2), Inches(2.66), Inches(2.7), Inches(0.24), h, size=18, bold=True)
        text(s, x + Inches(0.2), Inches(3.16), Inches(2.7), Inches(0.92), b, size=11.5, color=MUTED)
    shape(s, Inches(0.78), Inches(6.0), Inches(11.76), Inches(0.42), BLUE_SOFT, line=BLUE_SOFT)
    text(s, Inches(1.04), Inches(6.14), Inches(11.2), Inches(0.16), "市场缺的不是“更多职业信息”，而是一位能持续理解用户、动态推演方向并推动行动的 AI 职业教练。", size=12.5, bold=True)
    footer(s, "结论：职业决策需要从一次性建议，升级为持续陪伴的智能系统。")

    # 4 solution
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    header(s, 4, "AI职业教练把职业探索拆成三件事：看清自己、找到方向、推进现实行动")
    shape(s, Inches(0.9), Inches(2.26), Inches(11.54), Inches(3.0), WHITE, line=LINE)
    steps = [
        ("01", "看清自己", "通过共情式对话复现工作场景，挖掘用户的天赋、动机、优势与卡点。"),
        ("02", "找到方向", "结合能力迁移与趋势变化，输出更动态、更个性化的职业方向判断。"),
        ("03", "推进行动", "把认知结果转化为现实路径、准备重点与下一步动作。"),
    ]
    for i, (n, h, b) in enumerate(steps):
        x = Inches(1.18 + 3.72 * i)
        text(s, x, Inches(2.64), Inches(0.46), Inches(0.2), n, size=22, bold=True, color=BLUE)
        text(s, x, Inches(3.08), Inches(2.7), Inches(0.24), h, size=18, bold=True)
        text(s, x, Inches(3.52), Inches(2.86), Inches(0.84), b, size=11.2, color=MUTED)
        if i < 2:
            line(s, x + Inches(2.88), Inches(3.18), x + Inches(3.28), Inches(3.18), color=LINE, width=2)
    footer(s, "产品价值不是一份测评结果，而是一套从理解到行动的连续链路。")

    # 5 journey
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    header(s, 5, "用户体验 Journey Map", "一条从“模糊困惑”到“现实判断”的职业决策链路。")
    chip(s, Inches(10.52), Inches(0.92), "Self-awareness → Career decision", fill=WHITE, line=LINE, color=MUTED)
    line(s, Inches(0.78), Inches(1.98), Inches(12.0), Inches(1.98), color=LINE, width=1.2)
    fills = [GOLD, PEACH, MINT, BLUE, LAV]
    labels = [("AWARENESS", "困惑进入"), ("DISCOVERY", "对话理解"), ("INSIGHT", "职业报告"), ("DECISION", "现实路径地图"), ("PROGRESS", "策略推进")]
    for i, ((en, cn), fill) in enumerate(zip(labels, fills)):
        arrow_stage(s, Inches(0.78 + 2.42 * i), Inches(2.36), Inches(2.12), fill, en, cn)
    cards = [
        ("01", "困惑进入", "用户不需要先知道岗位名，只需要先说出真实卡点。", ["困惑高", "目标远"]),
        ("02", "对话理解", "通过自然对话提炼偏好、能力、天赋与工作方式。", ["被听见", "开始聚焦"]),
        ("03", "职业报告", "把潜在优势和方向感总结成可被看见的职业画像。", ["被理解", "看见方向"]),
        ("04", "现实路径地图", "把画像翻译成现实职业方向、岗位差异和准备重点。", ["更笃定", "接近目标"]),
        ("05", "策略推进", "方向清楚后，再进入适合自己的准备节奏。", ["有策略", "可持续"]),
    ]
    for i, (num, title, body, tags) in enumerate(cards):
        x = Inches(0.78 + 2.42 * i)
        panel(s, x, Inches(4.06), Inches(2.12), Inches(2.28))
        text(s, x + Inches(0.18), Inches(4.26), Inches(0.32), Inches(0.16), num, size=10, bold=True, color=BLUE)
        text(s, x + Inches(0.18), Inches(4.52), Inches(1.7), Inches(0.22), title, size=16, bold=True)
        text(s, x + Inches(0.18), Inches(4.98), Inches(1.74), Inches(0.84), body, size=10.2, color=MUTED)
        chip(s, x + Inches(0.18), Inches(5.86), tags[0], fill=BG, line=BG, color=MUTED, size=8)
        chip(s, x + Inches(1.02), Inches(5.86), tags[1], fill=BG, line=BG, color=MUTED, size=8)
    footer(s, "Journey Map 是产品叙事的主轴：从感受进入，到职业判断，再到策略推进。")

    # 6 reality path map
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    header(s, 6, "现实路径地图", "把“我适合什么”翻译成现实职业方向、岗位差异与行动优先级。")
    panel(s, Inches(0.78), Inches(2.06), Inches(3.1), Inches(4.18), title="输入基础")
    bullets(s, Inches(1.04), Inches(2.66), Inches(2.45), [
        "用户的天赋、优势、偏好与动机",
        "可迁移能力与过往经历",
        "现实约束：时间、资源、风险承受度",
    ], gap=0.74)
    panel(s, Inches(4.12), Inches(2.06), Inches(4.1), Inches(4.18), title="路径翻译引擎", fill=BLUE_SOFT, line=BLUE_SOFT)
    bullets(s, Inches(4.38), Inches(2.66), Inches(3.45), [
        "把职业画像映射为可选方向簇",
        "区分不同岗位的核心差异与门槛",
        "判断“更适合 / 更可达 / 更值得准备”的顺序",
        "输出每条路径的准备重点与风险提示",
    ], gap=0.62)
    panel(s, Inches(8.46), Inches(2.06), Inches(4.08), Inches(4.18), title="输出结果")
    bullets(s, Inches(8.72), Inches(2.66), Inches(3.35), [
        "现实职业方向：下一步该往哪类机会靠近",
        "岗位差异：这些路径到底区别在哪里",
        "准备重点：先补什么、怎么补、先做哪一步",
    ], gap=0.74)
    line(s, Inches(3.9), Inches(4.12), Inches(4.12), Inches(4.12), color=BLUE, width=2)
    line(s, Inches(8.24), Inches(4.12), Inches(8.46), Inches(4.12), color=BLUE, width=2)
    shape(s, Inches(4.54), Inches(5.7), Inches(3.26), Inches(0.42), PEACH_SOFT, line=PEACH_SOFT)
    text(s, Inches(4.78), Inches(5.84), Inches(2.78), Inches(0.14), "核心价值：让“方向感”变成“现实判断”。", size=10.5, bold=True)
    footer(s, "现实路径地图是产品的关键差异页，建议在路演时重点讲解。")

    # 7 product experience
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    header(s, 7, "产品体验从一段对话开始，而不是一张静态测评表")
    panel(s, Inches(0.78), Inches(2.06), Inches(4.08), Inches(4.26), fill=PEACH_SOFT, line=PEACH_SOFT)
    chip(s, Inches(1.04), Inches(2.34), "体验逻辑", fill=WHITE, line=WHITE, color=PEACH)
    bullets(s, Inches(1.04), Inches(2.84), Inches(3.34), [
        "以教练式提问帮助用户回到真实工作场景",
        "把“我适合什么工作”拆解成可追问、可验证的过程",
        "相比一次性测试，更适合高频迭代和持续校准",
        "更容易承接情绪、认知与行动层面的连续变化",
    ], gap=0.62)
    panel(s, Inches(5.4), Inches(2.06), Inches(6.82), Inches(4.26))
    picture(s, "image27.png", Inches(5.72), Inches(2.26), width=Inches(6.18))
    footer(s, "真实素材来源：对话理解界面。")

    # 8 product output
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    header(s, 8, "产品输出的价值，不是“给答案”，而是形成一份更可行动的职业分析")
    panel(s, Inches(0.78), Inches(2.06), Inches(6.26), Inches(4.28))
    picture(s, "image26.png", Inches(1.04), Inches(2.32), width=Inches(5.74))
    panel(s, Inches(7.34), Inches(2.06), Inches(5.2), Inches(4.28), title="报告核心输出", fill=MINT_SOFT, line=MINT_SOFT)
    bullets(s, Inches(7.62), Inches(2.72), Inches(4.3), [
        "个人关键特质总结：提炼能力、目标、愿景与优势",
        "天赋适配岗位：结合趋势给出更长期的发展方向建议",
        "可复用能力：识别可迁移、可跨岗位使用的竞争力",
        "岗位推荐与对比：帮助用户做更清晰的决策",
    ], gap=0.74)
    footer(s, "真实素材来源：职业镜像报告界面。")

    # 9 engine
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    header(s, 9, "产品底层由三层能力组成：对话理解、职业匹配、行动推进")
    panel(s, Inches(0.96), Inches(2.16), Inches(11.36), Inches(3.76))
    rows = [
        ("自然对话层", "使用多模型调试教练式对话，强调自然、共情与个性化理解。", PEACH),
        ("匹配分析层", "结合用户经历、天赋、性格与发展偏好，给出更适合的职业方向与岗位簇。", BLUE),
        ("行动推进层", "通过现实路径地图、准备重点、策略推进机制帮助用户真正开始行动。", MINT),
    ]
    for i, (h, b, c) in enumerate(rows):
        y = Inches(2.52 + 1.02 * i)
        shape(s, Inches(1.22), y, Inches(2.16), Inches(0.7), c, rounded=False)
        text(s, Inches(1.38), y + Inches(0.22), Inches(1.84), Inches(0.18), h, size=15, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        text(s, Inches(3.68), y + Inches(0.12), Inches(7.95), Inches(0.34), b, size=11.8, color=MUTED)
    footer(s, "当前材料已明确提到多模型参与训练与调试；此处保留为能力说明，不扩展未经验证的数据。")

    # 10 users
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    header(s, 10, "优先用户不是所有人，而是三类最容易感到职业失配的人群")
    personas = [
        ("应届生与毕业1-3年人群", ["需要更快理解自己适合什么", "缺少真实工作反馈与方向判断", "信息很多，但决策信心弱"], BLUE_SOFT),
        ("转岗/求职中的职场人", ["已有履历，但对下一步路径犹豫", "担心经验在新环境中失效", "需要把能力迁移成新的职业叙事"], PEACH_SOFT),
        ("高敏感、高思考型探索者", ["容易自我怀疑或长期停留在思考中", "需要共情式陪伴与行动推进", "更适合教练式而不是模板化建议"], MINT_SOFT),
    ]
    for i, (head, arr, fill) in enumerate(personas):
        x = Inches(0.78 + 4.0 * i)
        panel(s, x, Inches(2.08), Inches(3.76), Inches(4.12), title=head, fill=fill, line=fill)
        bullets(s, x + Inches(0.22), Inches(2.8), Inches(3.16), arr, gap=0.72)
    footer(s, "优先人群页面帮助解释产品为什么先从高需求、高焦虑、高决策密度人群切入。")

    # 11 business model
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    header(s, 11, "商业模式从轻量试用切入，再逐步扩展到订阅、陪跑与B端场景")
    models = [
        ("C端轻量订阅", ["基础对话、职业镜像报告、岗位推荐", "适合低门槛获客与持续使用", "验证用户愿不愿意为持续陪伴付费"], WHITE),
        ("高客单陪跑服务", ["围绕求职转化、转岗推进、行动落地", "AI + 人工辅导可形成更高价值组合", "适合验证高意愿用户的付费深度"], BLUE_SOFT),
        ("B端合作场景", ["高校就业指导、职业教育机构、人才服务平台", "可作为评估、辅导、转化工具的底层能力", "适合扩展规模化收入来源"], WHITE),
    ]
    for i, (head, arr, fill) in enumerate(models):
        x = Inches(0.78 + 4.0 * i)
        panel(s, x, Inches(2.1), Inches(3.76), Inches(4.1), title=head, fill=fill, line=LINE if fill == WHITE else fill)
        bullets(s, x + Inches(0.22), Inches(2.82), Inches(3.1), arr, gap=0.72)
    footer(s, "商业模式为合理推演，建议后续结合真实用户反馈和试运行数据进一步收敛。")

    # 12 market & moat
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    header(s, 12, "市场机会来自三个交叉趋势，而长期壁垒来自方法论、交互与持续沉淀")
    panel(s, Inches(0.78), Inches(2.06), Inches(5.46), Inches(4.3), title="为什么现在值得做")
    bullets(s, Inches(1.02), Inches(2.78), Inches(4.8), [
        "岗位变化更快，职业焦虑与转岗需求持续增加",
        "大模型让“高频陪伴式辅导”首次具备规模化与可负担性",
        "用户越来越难接受模板化建议，更愿意为被理解和可行动结果付费",
    ], gap=0.9)
    panel(s, Inches(6.72), Inches(2.06), Inches(5.82), Inches(4.3), title="长期壁垒来自什么", fill=PEACH_SOFT, line=PEACH_SOFT)
    bullets(s, Inches(6.98), Inches(2.78), Inches(5.1), [
        "方法论壁垒：把心理学、教练技术与职业规划逻辑结构化进产品",
        "交互壁垒：沉淀更适合职业场景的提问路径、反馈方式与行动编排",
        "数据壁垒：随着对话、判断与行动反馈积累，产品会越来越理解用户",
        "品牌壁垒：率先占领“AI职业教练”心智，会形成更强用户信任",
    ], gap=0.7)
    footer(s, "这页强调“为何值得做”与“长期如何站稳”，而不是堆砌未经验证的大数字。")

    # 13 competition
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    header(s, 13, "竞争不只是“谁也做职业规划”，关键在于谁能把理解与行动真正连起来")
    panel(s, Inches(0.78), Inches(2.06), Inches(11.76), Inches(4.48))
    xs = [0.98, 3.38, 6.18, 9.02]
    heads = ["方案类型", "典型优势", "常见短板", "AI职业教练的差异点"]
    for x, h in zip(xs, heads):
        text(s, Inches(x), Inches(2.34), Inches(1.9), Inches(0.16), h, size=10, bold=True, color=MUTED)
    rows = [
        ("传统职业咨询", "有真人经验与深度沟通", "价格高、频次低、难规模化", "更高频、更低摩擦，保留教练式交互体验"),
        ("测评/报告工具", "标准化强、交付快", "静态、模板化，难承接复杂个体差异", "从一次性结果升级为持续理解与校准"),
        ("通用AI聊天工具", "交互自然、门槛低", "缺少职业方法论与行动设计", "聚焦职业场景与现实路径，不止是聊天"),
    ]
    y = 2.78
    for row in rows:
        for x, val in zip(xs, row):
            text(s, Inches(x), Inches(y), Inches(2.06 if x < 9 else 2.6), Inches(0.7), val, size=11, color=INK if x == 0.98 else MUTED)
        shape(s, Inches(0.96), Inches(y + 0.82), Inches(11.2), Inches(0.012), LINE, rounded=False)
        y += 1.06
    footer(s, "竞争分析采用“类别对比”写法，避免无证据点名对标。")

    # 14 gtm
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    header(s, 14, "增长路径建议先做高相关场景切入，再逐步沉淀口碑、内容与合作分发")
    panel(s, Inches(0.92), Inches(2.2), Inches(11.2), Inches(3.8), fill=WHITE, line=LINE)
    steps = [
        ("1. 场景切入", "围绕求职、转岗、职业迷茫等强需求场景启动"),
        ("2. 内容获客", "通过真实案例与职业洞察内容建立认知与信任"),
        ("3. 体验转化", "以低门槛试用进入，再通过报告与行动工具提升转化"),
        ("4. 合作扩展", "向高校、教育机构、人才平台等B端合作场景延展"),
    ]
    for i, (h, b) in enumerate(steps):
        x = Inches(1.18 + 2.7 * i)
        text(s, x, Inches(2.74), Inches(0.34), Inches(0.18), f"{i+1:02d}", size=21, bold=True, color=BLUE)
        text(s, x, Inches(3.16), Inches(2.2), Inches(0.2), h, size=16, bold=True)
        text(s, x, Inches(3.56), Inches(2.18), Inches(0.82), b, size=10.8, color=MUTED)
        if i < 3:
            line(s, x + Inches(2.18), Inches(3.0), x + Inches(2.52), Inches(3.0), color=LINE, width=2)
    footer(s, "增长页为策略建议，后续可结合真实渠道数据进一步细化。")

    # 15 progress roadmap
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    header(s, 15, "目前已完成 MVP 核心链路，接下来重点是验证转化、打磨产品并建立增长飞轮")
    cols = [
        ("已完成", ["核心教练对话功能已初步实现", "职业镜像报告与岗位推荐功能已打通", "产品可通过网页链接进行体验"], WHITE),
        ("未来12个月重点", ["强化现实路径地图与行动推进能力", "优化个性化判断质量与推荐逻辑", "完成初步商业化与增长路径验证"], BLUE_SOFT),
        ("当前需要的资源", ["AI模型训练与推荐算法指导", "心理学与职业教练方法论支持", "商业化落地、品牌建设与用户增长协助"], WHITE),
    ]
    for i, (h, arr, fill) in enumerate(cols):
        x = Inches(0.78 + 4.0 * i)
        panel(s, x, Inches(2.1), Inches(3.76), Inches(4.12), title=h, fill=fill, line=LINE if fill == WHITE else fill)
        bullets(s, x + Inches(0.22), Inches(2.82), Inches(3.12), arr, gap=0.72)
    footer(s, "项目进展页承接融资诉求：说明目前已到哪个阶段，以及钱要用来解决什么问题。")

    # 16 funding use
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    header(s, 16, "融资诉求与资金用途", "当前更适合寻求天使阶段支持，用于把方法论产品化、把MVP推向验证期。")
    panel(s, Inches(0.78), Inches(2.02), Inches(4.02), Inches(4.34), title="建议融资用途", fill=PEACH_SOFT, line=PEACH_SOFT)
    bullets(s, Inches(1.04), Inches(2.74), Inches(3.36), [
        "产品与算法打磨：提升职业判断质量、路径地图与推荐逻辑",
        "用户研究与验证：跑通核心人群的留存、反馈与转化",
        "增长试验：验证内容获客、试用转化与合作分发效率",
        "基础团队补位：产品、AI算法、内容/增长方向的关键协作",
    ], gap=0.7)
    panel(s, Inches(5.12), Inches(2.02), Inches(3.14), Inches(4.34), title="建议比例")
    ratios = [("产品/算法", "40%"), ("用户研究/验证", "25%"), ("增长/品牌试验", "20%"), ("组织与运营", "15%")]
    y = 2.84
    for name, pct in ratios:
        shape(s, Inches(5.38), Inches(y), Inches(2.58), Inches(0.58), WHITE, line=LINE)
        text(s, Inches(5.62), Inches(y + 0.14), Inches(1.4), Inches(0.18), name, size=11.5, color=MUTED)
        text(s, Inches(7.1), Inches(y + 0.1), Inches(0.62), Inches(0.18), pct, size=18, bold=True, color=BLUE, align=PP_ALIGN.RIGHT)
        y += 0.82
    panel(s, Inches(8.58), Inches(2.02), Inches(3.96), Inches(4.34), title="12个月目标", fill=BLUE_SOFT, line=BLUE_SOFT)
    bullets(s, Inches(8.84), Inches(2.74), Inches(3.3), [
        "完成核心人群产品闭环验证",
        "形成更稳定的现实路径地图与行动推进机制",
        "明确 1-2 条可持续商业化路径",
        "沉淀可复制的增长与合作打法",
    ], gap=0.74)
    footer(s, "未虚构具体融资金额；如果你有目标金额，我可以继续补成更完整的融资页。")

    # 17 founder
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    header(s, 17, "项目由创始人独立推进，已具备明确的问题意识、方法论基础与产品原型")
    panel(s, Inches(0.78), Inches(2.08), Inches(7.18), Inches(4.18))
    text(s, Inches(1.04), Inches(2.42), Inches(2.0), Inches(0.2), "创始人：王旭", size=18, bold=True)
    bullets(s, Inches(1.04), Inches(2.96), Inches(6.38), [
        "目前主要由创始人负责产品设计、模型训练和代码研发",
        "中科院心理所心理学与人工智能 2018 级在职研究生",
        "8 年在线 IT 职业教育课程设计与教研经历",
        "对职业发展、学习设计与行动转化场景有长期积累",
    ], size=11.5, gap=0.72)
    panel(s, Inches(8.28), Inches(2.08), Inches(4.26), Inches(4.18), fill=MINT_SOFT, line=MINT_SOFT)
    text(s, Inches(8.56), Inches(2.42), Inches(3.2), Inches(0.2), "试用信息", size=17, bold=True)
    text(s, Inches(8.56), Inches(2.9), Inches(3.2), Inches(0.66), "产品链接：\nhttps://echotalent.fun/", size=11.4, color=MUTED)
    text(s, Inches(8.56), Inches(3.82), Inches(3.2), Inches(0.66), "说明：\n手机与电脑端均可使用，海外服务器可能偶有不稳定。", size=11.4, color=MUTED)
    text(s, Inches(8.56), Inches(4.84), Inches(3.2), Inches(0.66), "适合展示：\n对话理解、职业镜像报告、现实路径地图链路。", size=11.4, color=MUTED)
    footer(s, "如果需要，我可以继续补“融资金额 / 出让比例 / 投后目标”的标准投资人页。")

    prs.save(OUTPUT_PPTX)


if __name__ == "__main__":
    build()
