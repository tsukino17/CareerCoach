from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import MSO_VERTICAL_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


BASE_DIR = Path("/Users/sue/Documents/trae_projects/career/outputs/manual-20260513-115102/presentations/ai-career-coach-bp")
ASSET_DIR = BASE_DIR / "template-inspect" / "assets" / "ppt" / "media"
OUTPUT_PPTX = BASE_DIR / "output" / "AI职业教练-商业BP-v2-可编辑版.pptx"

NAVY = RGBColor(26, 34, 56)
BLUE = RGBColor(77, 106, 245)
SKY = RGBColor(100, 173, 255)
MINT = RGBColor(99, 194, 165)
PEACH = RGBColor(247, 168, 132)
INK = RGBColor(35, 42, 58)
MUTED = RGBColor(104, 113, 133)
SOFT = RGBColor(241, 244, 250)
WHITE = RGBColor(255, 255, 255)
LINE = RGBColor(222, 228, 236)
PALE_BLUE = RGBColor(235, 242, 255)
PALE_MINT = RGBColor(232, 247, 242)
PALE_PEACH = RGBColor(255, 241, 234)


def bg(slide, color=WHITE):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def rect(slide, left, top, width, height, color, line=None, radius=False):
    shape_type = MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    shp = slide.shapes.add_shape(shape_type, left, top, width, height)
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    if radius:
        shp.adjustments[0] = 0.08
    if line:
        shp.line.color.rgb = line
        shp.line.width = Pt(1)
    else:
        shp.line.fill.background()
    return shp


def text(slide, left, top, width, height, value, *, size=18, bold=False, color=INK,
         align=PP_ALIGN.LEFT, font_name="PingFang SC", line_spacing=1.2):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = MSO_VERTICAL_ANCHOR.TOP
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = line_spacing
    run = p.add_run()
    run.text = value
    font = run.font
    font.name = font_name
    font.size = Pt(size)
    font.bold = bold
    font.color.rgb = color
    return box


def kicker(slide, page, label="AI CAREER COACH"):
    text(slide, Inches(0.72), Inches(0.38), Inches(2.2), Inches(0.2), label, size=9, bold=True, color=MUTED)
    text(slide, Inches(12.0), Inches(0.38), Inches(0.45), Inches(0.2), f"{page:02d}", size=9, color=MUTED, align=PP_ALIGN.RIGHT)


def title_block(slide, title, subtitle=None, width=8.8):
    text(slide, Inches(0.72), Inches(0.82), Inches(width), Inches(0.72), title, size=28, bold=True)
    if subtitle:
        text(slide, Inches(0.72), Inches(1.55), Inches(10.6), Inches(0.42), subtitle, size=11, color=MUTED)


def footer(slide, value):
    rect(slide, Inches(0.72), Inches(6.7), Inches(11.15), Inches(0.012), LINE)
    text(slide, Inches(0.72), Inches(6.78), Inches(11.0), Inches(0.16), value, size=8, color=MUTED)


def chip(slide, left, top, value, fill=WHITE, line=LINE, color=MUTED):
    w = max(0.7, 0.14 * len(value) + 0.42)
    rect(slide, left, top, Inches(w), Inches(0.28), fill, line=line, radius=True)
    text(slide, left + Inches(0.12), top + Inches(0.05), Inches(w - 0.24), Inches(0.14), value, size=8.5, bold=True, color=color)


def metric(slide, left, top, value, label, note, accent=BLUE):
    rect(slide, left, top, Inches(0.03), Inches(0.72), accent)
    text(slide, left + Inches(0.12), top - Inches(0.02), Inches(1.7), Inches(0.3), value, size=24, bold=True)
    text(slide, left + Inches(0.12), top + Inches(0.31), Inches(1.9), Inches(0.18), label, size=9, bold=True, color=MUTED)
    text(slide, left + Inches(0.12), top + Inches(0.49), Inches(2.1), Inches(0.18), note, size=8, color=MUTED)


def panel(slide, left, top, width, height, title=None, fill=WHITE, line=LINE, title_color=INK):
    rect(slide, left, top, width, height, fill, line=line, radius=True)
    if title:
        text(slide, left + Inches(0.2), top + Inches(0.18), width - Inches(0.4), Inches(0.28), title, size=16, bold=True, color=title_color)


def bullets(slide, left, top, width, items, *, size=11, color=MUTED, gap=0.44):
    y = top
    for item in items:
        text(slide, left, y, width, Inches(0.36), f"• {item}", size=size, color=color)
        y += Inches(gap)


def add_image(slide, filename, left, top, *, width=None, height=None):
    path = ASSET_DIR / filename
    if width and height:
        slide.shapes.add_picture(str(path), left, top, width=width, height=height)
    elif width:
        slide.shapes.add_picture(str(path), left, top, width=width)
    elif height:
        slide.shapes.add_picture(str(path), left, top, height=height)
    else:
        slide.shapes.add_picture(str(path), left, top)


def connector(slide, x1, y1, x2, y2, color=BLUE, width=1.5):
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    line.line.color.rgb = color
    line.line.width = Pt(width)
    return line


def journey_arrow(slide, left, top, width, height, color, en, cn):
    shp = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.CHEVRON, left, top, width, height)
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    shp.line.fill.background()
    text(slide, left + Inches(0.15), top + Inches(0.23), width - Inches(0.3), Inches(0.18), en, size=10, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    text(slide, left + Inches(0.15), top + Inches(0.55), width - Inches(0.3), Inches(0.28), cn, size=18, bold=True, color=WHITE, align=PP_ALIGN.CENTER)


def mini_card(slide, left, top, num, title, body, tags):
    panel(slide, left, top, Inches(2.32), Inches(2.7), fill=WHITE, line=LINE)
    rect(slide, left + Inches(0.16), top + Inches(0.33), Inches(0.46), Inches(0.46), SOFT, radius=True)
    text(slide, left + Inches(0.83), top + Inches(0.36), Inches(0.34), Inches(0.2), num, size=10, bold=True, color=BLUE)
    text(slide, left + Inches(0.83), top + Inches(0.6), Inches(1.2), Inches(0.24), title, size=17, bold=True)
    text(slide, left + Inches(0.18), top + Inches(1.1), Inches(1.95), Inches(0.92), body, size=10.5, color=MUTED)
    x = left + Inches(0.18)
    y = top + Inches(2.12)
    for tag in tags:
        chip(slide, x, y, tag, fill=SOFT, line=SOFT, color=MUTED)
        x += Inches(0.72)


def build():
    OUTPUT_PPTX.parent.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # 1 cover
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    rect(s, Inches(7.85), Inches(0), Inches(5.48), Inches(7.5), SOFT)
    rect(s, Inches(7.85), Inches(0), Inches(5.48), Inches(2.2), PALE_BLUE)
    rect(s, Inches(0.74), Inches(1.56), Inches(0.08), Inches(1.12), BLUE)
    text(s, Inches(1.0), Inches(1.48), Inches(5.1), Inches(0.6), "AI职业教练", size=32, bold=True, color=BLUE)
    text(s, Inches(1.0), Inches(2.16), Inches(6.0), Inches(0.9), "帮助用户在AI时代完成职业定位、能力迁移与行动转化", size=24, bold=True)
    text(s, Inches(1.0), Inches(3.14), Inches(5.0), Inches(1.0), "一款结合教练式对话、职业分析报告与行动力工具箱的 AI 职业发展产品。", size=13, color=MUTED)
    metric(s, Inches(1.0), Inches(5.12), "MVP", "当前阶段", "核心链路已初步打通")
    metric(s, Inches(3.22), Inches(5.12), "3层能力", "产品结构", "对话 / 匹配 / 行动")
    metric(s, Inches(5.44), Inches(5.12), "可编辑BP", "本次交付", "已补齐关键路演缺项")
    panel(s, Inches(8.38), Inches(1.0), Inches(4.2), Inches(5.35), fill=WHITE, line=LINE)
    add_image(s, "image26.png", Inches(8.7), Inches(1.35), width=Inches(3.56))
    chip(s, Inches(8.72), Inches(5.35), "真实产品报告界面", fill=PALE_PEACH, line=PALE_PEACH, color=PEACH)
    text(s, Inches(8.72), Inches(5.72), Inches(3.45), Inches(0.45), "不是停留在概念展示，而是已经能够给出职业镜像报告与结构化洞察。", size=10.5, color=MUTED)
    footer(s, "基于原始《产品核心介绍》扩展为更完整的商业 BP。")

    # 2 opportunity
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    kicker(s, 2)
    title_block(s, "AI正在重构职业发展决策，用户需要的不是信息更多，而是判断更准、行动更快")
    panel(s, Inches(0.72), Inches(2.05), Inches(3.82), Inches(3.78), title="岗位与技能失效更快")
    bullets(s, Inches(0.95), Inches(2.58), Inches(3.3), [
        "岗位半衰期缩短，热点赛道与能力要求持续变化",
        "单一专业与固定履历不再天然构成长期竞争力",
        "用户更需要识别可迁移、可复用的底层能力"
    ])
    panel(s, Inches(4.76), Inches(2.05), Inches(3.82), Inches(3.78), title="传统职业规划明显滞后", fill=PALE_BLUE, line=PALE_BLUE)
    bullets(s, Inches(4.99), Inches(2.58), Inches(3.3), [
        "高校课程与真实岗位要求进一步脱节",
        "转岗、求职、重返市场的人群更容易感到失配",
        "用户需要更贴近现实场景的动态判断系统"
    ])
    panel(s, Inches(8.8), Inches(2.05), Inches(3.82), Inches(3.78), title="AI既带来机会，也放大焦虑")
    bullets(s, Inches(9.03), Inches(2.58), Inches(3.3), [
        "企业更强调轻量化、效率与结果导向",
        "个人一边看到新机会，一边感到被替代压力",
        "职业辅导不止要给方向，更要帮助用户开始行动"
    ])
    footer(s, "原始材料来源：时代背景、用户痛点页。")

    # 3 gap
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s, SOFT)
    kicker(s, 3)
    title_block(s, "当前职业辅导产品的问题，不是没有内容，而是太慢、太贵、太模板化")
    panel(s, Inches(0.72), Inches(2.02), Inches(11.9), Inches(3.85), fill=WHITE, line=LINE)
    for i, (head, body, accent) in enumerate([
        ("流程繁琐", "填表、测评、上传简历、等待分析，体验链路长，反馈时效慢。", BLUE),
        ("成本高、频次低", "真人咨询通常价格不低，且多为按周预约，难承接高频决策。", PEACH),
        ("推荐静态且同质化", "容易围绕热门岗位给出答案，忽略个体差异与长期潜力。", MINT),
    ]):
        x = Inches(1.0 + 3.95 * i)
        rect(s, x, Inches(2.4), Inches(0.08), Inches(0.95), accent)
        text(s, x + Inches(0.18), Inches(2.42), Inches(2.8), Inches(0.28), head, size=18, bold=True)
        text(s, x + Inches(0.18), Inches(2.92), Inches(3.0), Inches(1.0), body, size=12, color=MUTED)
    panel(s, Inches(0.72), Inches(6.0), Inches(11.9), Inches(0.5), fill=PALE_PEACH, line=PALE_PEACH)
    text(s, Inches(0.95), Inches(6.16), Inches(11.3), Inches(0.2), "市场缺的不是“更多职业信息”，而是一位能持续理解用户、动态推演方向并推动行动的 AI 职业教练。", size=13, bold=True, color=INK)
    footer(s, "结论：职业决策需要从一次性建议，升级为持续陪伴的智能系统。")

    # 4 product
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    kicker(s, 4)
    title_block(s, "AI职业教练把职业探索拆成三件事：看清自己、找到方向、推进现实行动")
    rect(s, Inches(0.78), Inches(2.18), Inches(11.76), Inches(3.14), PALE_BLUE, line=PALE_BLUE, radius=True)
    blocks = [
        ("01", "看清自己", "通过共情式对话复现真实工作场景，挖掘用户的天赋、动机、优势与卡点。"),
        ("02", "找到方向", "结合能力迁移与时代趋势，给出更动态、更个性化的岗位与路径建议。"),
        ("03", "推进行动", "通过提示、记录、演练与模拟，帮助用户把认知结果转化为现实行动。"),
    ]
    for i, (n, h, b) in enumerate(blocks):
        x = Inches(1.08 + 3.9 * i)
        text(s, x, Inches(2.56), Inches(0.6), Inches(0.3), n, size=23, bold=True, color=BLUE)
        text(s, x, Inches(3.0), Inches(2.7), Inches(0.28), h, size=18, bold=True)
        text(s, x, Inches(3.45), Inches(2.95), Inches(1.0), b, size=11, color=MUTED)
        if i < 2:
            connector(s, x + Inches(2.95), Inches(3.13), x + Inches(3.42), Inches(3.13))
    text(s, Inches(0.78), Inches(5.75), Inches(11.5), Inches(0.24), "对用户而言，这不是一张测评结果，而是一位会持续理解、校准并推动落地的 AI 教练。", size=12, color=MUTED)
    footer(s, "结构重组自：产品理念、职业分析报告、岗位推荐与行动力工具箱页。")

    # 5 principles
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    kicker(s, 5)
    title_block(s, "产品核心判断：在快速变化的时代，真正稳定的是能力、动机与天赋")
    for i, (head, body, fill) in enumerate([
        ("能力与特质", "相比只问“会不会某个工具”，更重要的是用户是否具备可迁移的审美、分析、表达、组织与落地能力。", WHITE),
        ("动机与热情", "职业选择背后有更深层的驱动力。看见用户真正想做什么，才能找到更宽的职业边界。", PALE_PEACH),
        ("天赋与特长", "长期优势往往来自那些做起来最顺手、最有正反馈的事情，而不只是短期热门岗位。", WHITE),
    ]):
        x = Inches(0.72 + 4.05 * i)
        panel(s, x, Inches(2.08), Inches(3.72), Inches(3.4), title=head, fill=fill, line=LINE)
        text(s, x + Inches(0.22), Inches(2.7), Inches(3.2), Inches(1.55), body, size=11, color=MUTED)
    panel(s, Inches(0.72), Inches(5.88), Inches(11.9), Inches(0.48), fill=PALE_BLUE, line=PALE_BLUE)
    text(s, Inches(0.98), Inches(6.03), Inches(11.2), Inches(0.2), "因此，AI职业教练的价值不只是“推荐岗位”，而是帮助用户识别自己的长期竞争力，并据此做出更好的职业决策。", size=12.5, bold=True)
    footer(s, "原始材料来源：快速变化的AI时代什么是不变的。")

    # 6 methodology
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s, SOFT)
    kicker(s, 6)
    title_block(s, "方法论不是装饰，而是产品可信度的一部分")
    items = [
        ("福格行为模型", "行为 = 动机 + 能力 + 提示，用更低摩擦的方式推动用户开始行动。", PALE_BLUE),
        ("TRUTH转化模型", "围绕觉察、转化、行动，帮助用户把职业兴趣逐步落到现实。", PALE_PEACH),
        ("教练式对话", "不是直接告诉用户答案，而是通过启发式追问帮助其看见资源与渴望。", WHITE),
        ("焦点解决短程疗法", "帮助迷茫焦虑的用户重新聚焦已有优势，建立可行动的信心。", PALE_MINT),
    ]
    for idx, (head, body, fill) in enumerate(items):
        x = Inches(0.72 + (idx % 2) * 5.9)
        y = Inches(2.1 + (idx // 2) * 1.95)
        panel(s, x, y, Inches(5.7), Inches(1.56), title=head, fill=fill, line=LINE)
        text(s, x + Inches(0.22), y + Inches(0.58), Inches(5.1), Inches(0.64), body, size=11, color=MUTED)
    footer(s, "原始材料来源：产品理论与技术支持页。")

    # 7 engine
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    kicker(s, 7)
    title_block(s, "产品底层由三层能力组成：对话理解、职业匹配、行动转化")
    panel(s, Inches(0.72), Inches(2.08), Inches(11.9), Inches(3.9), fill=WHITE, line=LINE)
    rows = [
        ("自然对话层", "使用多模型调试教练式对话，强调自然、共情与个性化理解。", BLUE),
        ("匹配分析层", "结合用户经历、天赋、性格与发展偏好，给出更适合的岗位、公司与行业建议。", PEACH),
        ("行动工具层", "通过提示、行动日记、方案演练、模拟面试等工具帮助用户推进现实转化。", MINT),
    ]
    for i, (head, body, accent) in enumerate(rows):
        y = Inches(2.48 + 1.1 * i)
        rect(s, Inches(1.02), y, Inches(2.05), Inches(0.78), accent)
        text(s, Inches(1.27), y + Inches(0.22), Inches(1.55), Inches(0.2), head, size=15, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        text(s, Inches(3.33), y + Inches(0.15), Inches(8.72), Inches(0.48), body, size=12, color=MUTED)
    text(s, Inches(0.72), Inches(6.15), Inches(11.2), Inches(0.22), "当前材料已明确提到使用 qwen-plus、gemini、deepseek 等模型进行训练与调试；这里保留为能力说明，不扩展未经验证的性能数据。", size=9.5, color=MUTED)
    footer(s, "原始材料来源：AI智能底层驱动页。")

    # 8 experience
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    kicker(s, 8)
    title_block(s, "产品体验从一段对话开始，而不是一张静态测评表")
    panel(s, Inches(0.72), Inches(2.06), Inches(4.2), Inches(4.38), fill=PALE_BLUE, line=PALE_BLUE)
    chip(s, Inches(0.98), Inches(2.34), "体验逻辑", fill=WHITE, line=WHITE, color=BLUE)
    bullets(s, Inches(0.98), Inches(2.84), Inches(3.45), [
        "以教练式提问帮助用户回到真实工作场景",
        "把“我适合什么工作”拆解成可追问、可验证的过程",
        "相比一次性测试，更适合高频迭代和持续校准",
        "更容易承接情绪、认知与行动层面的连续变化"
    ], size=11.2, color=MUTED, gap=0.58)
    add_image(s, "image27.png", Inches(6.0), Inches(1.68), width=Inches(6.0))
    footer(s, "真实素材来源：原始PPT中的产品界面演示页。")

    # 9 output
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s, SOFT)
    kicker(s, 9)
    title_block(s, "产品输出的价值，不是“给答案”，而是形成一份更可行动的职业分析")
    add_image(s, "image26.png", Inches(0.82), Inches(1.95), width=Inches(6.2))
    panel(s, Inches(7.36), Inches(2.06), Inches(5.1), Inches(4.18), title="报告核心输出", fill=WHITE, line=LINE)
    bullets(s, Inches(7.62), Inches(2.66), Inches(4.45), [
        "个人关键特质总结：提炼能力、目标、愿景与优势",
        "天赋适配岗位：结合趋势给出更长期的发展方向建议",
        "可复用能力：识别可迁移、可跨岗位使用的竞争力",
        "岗位推荐与对比：帮助用户做更清晰的决策",
    ], gap=0.68)
    panel(s, Inches(7.62), Inches(5.54), Inches(3.9), Inches(0.48), fill=PALE_PEACH, line=PALE_PEACH)
    text(s, Inches(7.84), Inches(5.69), Inches(3.45), Inches(0.16), "重点是让用户理解“为什么适合”，而不是只看到岗位名称。", size=10.5, bold=True)
    footer(s, "真实素材来源：深度职业分析报告与职业镜像报告界面。")

    # 10 journey map
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    kicker(s, 10)
    title_block(s, "用户体验 Journey Map", "一条从“模糊困惑”到“现实判断”的职业决策链路。")
    chip(s, Inches(10.45), Inches(0.9), "Self-awareness → Career decision", fill=SOFT, line=SOFT, color=MUTED)
    rect(s, Inches(0.72), Inches(1.98), Inches(11.9), Inches(0.018), LINE)
    colors = [
        RGBColor(248, 172, 57),
        RGBColor(248, 91, 119),
        RGBColor(61, 196, 183),
        RGBColor(77, 149, 232),
        RGBColor(92, 91, 228),
    ]
    labels = [
        ("AWARENESS", "困惑进入"),
        ("DISCOVERY", "对话理解"),
        ("INSIGHT", "职业报告"),
        ("DECISION", "现实路径地图"),
        ("PROGRESS", "策略推进"),
    ]
    for i, ((en, cn), color) in enumerate(zip(labels, colors)):
        journey_arrow(s, Inches(0.72 + 2.48 * i), Inches(2.38), Inches(2.24), Inches(1.22), color, en, cn)
    mini_card(s, Inches(0.72), Inches(4.1), "01", "困惑进入", "用户不需要先知道岗位名，只需要先说出真实卡点。", ["困惑高", "目标远"])
    mini_card(s, Inches(3.16), Inches(4.1), "02", "对话理解", "通过自然对话提炼偏好、能力、天赋与工作方式。", ["被听见", "开始聚焦"])
    mini_card(s, Inches(5.6), Inches(4.1), "03", "职业报告", "把潜在优势和方向感总结成可被看见的画像。", ["被理解", "看见方向"])
    mini_card(s, Inches(8.04), Inches(4.1), "04", "现实路径地图", "把画像翻译成现实职业方向、岗位差异和准备重点。", ["更笃定", "接近目标"])
    mini_card(s, Inches(10.48), Inches(4.1), "05", "策略推进", "方向清楚后，再进入适合自己的准备节奏。", ["有策略", "可持续"])
    footer(s, "新增：Journey Map 用于清晰解释用户如何从感受进入，逐步走到现实职业决策。")

    # 11 reality path map
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s, SOFT)
    kicker(s, 11)
    title_block(s, "现实路径地图", "把“我适合什么”翻译成现实职业方向、岗位差异与行动优先级。")
    panel(s, Inches(0.72), Inches(2.02), Inches(3.0), Inches(4.3), title="输入基础", fill=WHITE, line=LINE)
    bullets(s, Inches(0.98), Inches(2.64), Inches(2.35), [
        "用户的天赋、优势、偏好与动机",
        "可迁移能力与过往经历",
        "现实约束：时间、资源、风险承受度"
    ], gap=0.72)
    panel(s, Inches(4.04), Inches(2.02), Inches(4.0), Inches(4.3), title="路径翻译引擎", fill=PALE_BLUE, line=PALE_BLUE)
    bullets(s, Inches(4.3), Inches(2.64), Inches(3.45), [
        "把职业画像映射为可选方向簇",
        "区分不同岗位的核心差异与门槛",
        "判断“更适合 / 更可达 / 更值得准备”的顺序",
        "输出每条路径的准备重点与风险提示"
    ], gap=0.62)
    panel(s, Inches(8.36), Inches(2.02), Inches(4.26), Inches(4.3), title="输出结果", fill=WHITE, line=LINE)
    bullets(s, Inches(8.62), Inches(2.64), Inches(3.6), [
        "现实职业方向：我下一步该往哪类机会靠近",
        "岗位差异：这些岗位究竟区别在哪里",
        "准备重点：先补什么、怎么补、先做哪一步"
    ], gap=0.72)
    connector(s, Inches(3.75), Inches(4.12), Inches(4.04), Inches(4.12), color=BLUE, width=2)
    connector(s, Inches(8.05), Inches(4.12), Inches(8.36), Inches(4.12), color=BLUE, width=2)
    panel(s, Inches(4.36), Inches(5.54), Inches(3.38), Inches(0.54), fill=PALE_PEACH, line=PALE_PEACH)
    text(s, Inches(4.6), Inches(5.71), Inches(2.9), Inches(0.18), "核心价值：让“方向感”变成“现实判断”。", size=10.5, bold=True)
    footer(s, "新增：现实路径地图页单独展开产品的关键决策价值。")

    # 12 users
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    kicker(s, 12)
    title_block(s, "优先用户不是所有人，而是三类最容易感到职业失配的人群")
    personas = [
        ("应届生与毕业1-3年人群", ["需要更快理解自己适合什么", "缺少真实工作反馈与方向判断", "对岗位信息多，但决策信心弱"], PALE_BLUE),
        ("转岗/求职中的职场人", ["已有履历，但对下一步路径犹豫", "担心过往经验在新环境中失效", "需要把能力迁移成新的职业叙事"], PALE_PEACH),
        ("高敏感、高思考型探索者", ["容易自我怀疑或长期停留在思考中", "需要共情式陪伴与行动推进", "更适合教练式而不是模板化建议"], PALE_MINT),
    ]
    for i, (head, arr, fill) in enumerate(personas):
        x = Inches(0.72 + 4.05 * i)
        panel(s, x, Inches(2.08), Inches(3.72), Inches(3.85), title=head, fill=fill, line=LINE)
        bullets(s, x + Inches(0.22), Inches(2.74), Inches(3.18), arr, size=11, color=MUTED, gap=0.62)
    footer(s, "用户分层为推演结果，用于BP表达，不代表已完成量化市场验证。")

    # 13 market
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    kicker(s, 13)
    title_block(s, "市场机会来自三个交叉趋势：职业焦虑上升、AI普及、个性化服务需求增长")
    panel(s, Inches(0.72), Inches(2.1), Inches(11.9), Inches(3.9), fill=WHITE, line=LINE)
    for i, (head, body, accent) in enumerate([
        ("职业重构", "岗位变化快、转岗频繁，让更多人需要更高频的职业判断服务。", BLUE),
        ("AI平权", "大模型让“高频陪伴式辅导”首次具备规模化与可负担性。", PEACH),
        ("个性化升级", "用户越来越难接受模板化建议，更愿意为深度理解和行动结果买单。", MINT),
    ]):
        x = Inches(1.0 + 3.93 * i)
        chip(s, x, Inches(2.42), head, fill=WHITE, line=accent, color=accent)
        text(s, x, Inches(2.95), Inches(3.0), Inches(1.0), body, size=12, color=MUTED)
    panel(s, Inches(0.72), Inches(6.1), Inches(11.9), Inches(0.38), fill=PALE_BLUE, line=PALE_BLUE)
    text(s, Inches(0.95), Inches(6.22), Inches(11.2), Inches(0.14), "这页强调的是“为什么现在值得做”，而不是在缺少外部研究支持时虚构 TAM/SAM/SOM 数字。", size=9.5, color=MUTED)
    footer(s, "市场页采用方向性表达，避免未经验证的定量夸张。")

    # 14 business model
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s, SOFT)
    kicker(s, 14)
    title_block(s, "商业模式可以从轻量试用切入，再逐步扩展到订阅、陪跑与B端场景")
    panel(s, Inches(0.72), Inches(2.08), Inches(3.7), Inches(4.16), title="C端轻量订阅", fill=WHITE, line=LINE)
    bullets(s, Inches(0.96), Inches(2.68), Inches(3.1), [
        "基础对话、职业镜像报告、岗位推荐",
        "适合低门槛获客与持续使用",
        "验证用户愿不愿意为持续陪伴付费"
    ], gap=0.62)
    panel(s, Inches(4.82), Inches(2.08), Inches(3.7), Inches(4.16), title="高客单陪跑服务", fill=WHITE, line=LINE)
    bullets(s, Inches(5.06), Inches(2.68), Inches(3.1), [
        "围绕求职转化、转岗推进、行动计划落地",
        "AI + 人工辅导可形成更高价值组合",
        "适合验证高意愿人群的付费深度"
    ], gap=0.62)
    panel(s, Inches(8.92), Inches(2.08), Inches(3.7), Inches(4.16), title="B端合作场景", fill=WHITE, line=LINE)
    bullets(s, Inches(9.16), Inches(2.68), Inches(3.1), [
        "高校就业指导、职业教育机构、人才服务平台",
        "可作为评估、辅导、转化工具的底层能力",
        "适合扩展品牌触达与规模化收入来源"
    ], gap=0.62)
    footer(s, "商业模式为现阶段合理推演，建议后续结合用户反馈和试运行数据进一步收敛。")

    # 15 competition
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    kicker(s, 15)
    title_block(s, "竞争不只是“谁也做职业规划”，关键在于谁能把理解与行动真正连起来")
    panel(s, Inches(0.72), Inches(2.06), Inches(11.9), Inches(4.3), fill=WHITE, line=LINE)
    headers = ["方案类型", "典型优势", "常见短板", "AI职业教练的差异点"]
    xs = [0.98, 3.35, 6.1, 9.0]
    for x, head in zip(xs, headers):
        text(s, Inches(x), Inches(2.34), Inches(1.8), Inches(0.22), head, size=10, bold=True, color=MUTED)
    rows = [
        ("传统职业咨询", "有真人经验与深度沟通", "价格高、频次低、难规模化", "更高频、更低摩擦，保留教练式交互体验"),
        ("测评/报告工具", "标准化强、交付快", "静态、模板化，难承接复杂个体差异", "从一次性结果升级为持续理解与校准"),
        ("通用AI聊天工具", "使用门槛低、交互自然", "缺少职业方法论与行动设计", "聚焦职业场景与行动转化，不止是聊天"),
    ]
    y = 2.78
    for row in rows:
        for x, value in zip(xs, row):
            text(s, Inches(x), Inches(y), Inches(2.05 if x < 9 else 2.6), Inches(0.8), value, size=11, color=INK if x == 0 else MUTED)
        rect(s, Inches(0.92), Inches(y + 0.88), Inches(11.2), Inches(0.012), LINE)
        y += 1.08
    footer(s, "竞争分析采用“类别对比”写法，避免无证据点名对标。")

    # 16 gtm
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    kicker(s, 16)
    title_block(s, "增长路径建议先做高相关场景切入，再逐步沉淀口碑、内容与合作分发")
    panel(s, Inches(0.72), Inches(2.18), Inches(11.9), Inches(3.92), fill=PALE_BLUE, line=PALE_BLUE)
    steps = [
        ("1. 场景切入", "围绕求职、转岗、职业迷茫等强需求场景启动"),
        ("2. 内容获客", "通过真实案例、职业洞察内容建立认知与信任"),
        ("3. 体验转化", "以低门槛试用进入，再通过报告与行动工具提升转化"),
        ("4. 合作扩展", "向高校、教育机构、人才平台等B端合作场景延展"),
    ]
    for i, (head, body) in enumerate(steps):
        x = Inches(1.0 + 2.9 * i)
        text(s, x, Inches(2.7), Inches(0.4), Inches(0.28), f"{i+1:02d}", size=22, bold=True, color=BLUE)
        text(s, x, Inches(3.16), Inches(2.4), Inches(0.24), head, size=16, bold=True)
        text(s, x, Inches(3.55), Inches(2.35), Inches(0.8), body, size=10.8, color=MUTED)
        if i < 3:
            connector(s, x + Inches(2.3), Inches(3.0), x + Inches(2.72), Inches(3.0), color=BLUE)
    footer(s, "增长页为策略建议，后续可结合真实渠道数据继续细化。")

    # 17 moat
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s, SOFT)
    kicker(s, 17)
    title_block(s, "长期壁垒不来自单一模型，而来自方法论、产品结构与持续交互数据的复合沉淀")
    panel(s, Inches(0.72), Inches(2.08), Inches(11.9), Inches(4.16), fill=WHITE, line=LINE)
    moat = [
        ("方法论壁垒", "把心理学、教练技术与职业规划逻辑结构化进产品。", BLUE),
        ("交互壁垒", "沉淀更适合职业场景的提问路径、反馈方式与行动编排。", PEACH),
        ("数据壁垒", "随着对话、判断与行动反馈积累，产品会越来越理解用户。", MINT),
        ("品牌壁垒", "如果率先占领“AI职业教练”心智，将形成更强的用户信任。", BLUE),
    ]
    for i, (head, body, accent) in enumerate(moat):
        x = Inches(1.0 + 2.78 * i)
        rect(s, x, Inches(2.58), Inches(0.08), Inches(2.58), accent)
        text(s, x + Inches(0.18), Inches(2.62), Inches(2.2), Inches(0.24), head, size=16, bold=True)
        text(s, x + Inches(0.18), Inches(3.1), Inches(2.2), Inches(1.1), body, size=11, color=MUTED)
    footer(s, "壁垒页强调中长期建设方向，不等同于已经形成的市场优势。")

    # 18 progress
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    kicker(s, 18)
    title_block(s, "目前已完成 MVP 核心链路，下一阶段重点是验证转化、打磨产品与建立增长飞轮")
    for i, (head, arr, fill) in enumerate([
        ("已完成", ["核心教练对话功能已初步实现", "职业分析报告与岗位推荐功能已打通", "产品可通过网页链接进行体验"], PALE_BLUE),
        ("下一阶段重点", ["加强行动力转化与持续陪跑能力", "优化推荐机制与个性化判断质量", "验证最适合的商业模式与增长路径"], PALE_PEACH),
        ("当前需要的资源", ["AI模型训练与推荐算法指导", "心理学与职业教练方法论支持", "商业化落地、品牌建设与用户增长协助"], PALE_MINT),
    ]):
        x = Inches(0.72 + 4.05 * i)
        panel(s, x, Inches(2.06), Inches(3.72), Inches(3.96), title=head, fill=fill, line=LINE)
        bullets(s, x + Inches(0.22), Inches(2.7), Inches(3.18), arr, gap=0.72)
    footer(s, "原始材料来源：项目进度、所需资源页。")

    # 19 founder
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    kicker(s, 19)
    title_block(s, "项目由创始人独立推进，已具备明确的问题意识、方法论基础与产品原型")
    panel(s, Inches(0.72), Inches(2.08), Inches(7.2), Inches(4.18), fill=WHITE, line=LINE)
    text(s, Inches(0.98), Inches(2.42), Inches(2.2), Inches(0.24), "创始人：王旭", size=18, bold=True)
    bullets(s, Inches(0.98), Inches(2.96), Inches(6.45), [
        "目前主要由创始人负责产品设计、模型训练和代码研发",
        "中科院心理所心理学与人工智能 2018 级在职研究生",
        "8 年在线 IT 职业教育课程设计与教研经历",
        "对职业发展、学习设计与行动转化场景有长期积累"
    ], size=11.5, gap=0.68)
    panel(s, Inches(8.2), Inches(2.08), Inches(4.42), Inches(4.18), fill=PALE_BLUE, line=PALE_BLUE)
    text(s, Inches(8.48), Inches(2.42), Inches(3.5), Inches(0.24), "试用信息", size=17, bold=True)
    text(s, Inches(8.48), Inches(2.9), Inches(3.35), Inches(0.72), "产品链接：\nhttps://echotalent.fun/", size=11.5, color=MUTED)
    text(s, Inches(8.48), Inches(3.82), Inches(3.35), Inches(0.72), "说明：\n手机与电脑端均可使用，海外服务器可能偶有不稳定。", size=11.5, color=MUTED)
    text(s, Inches(8.48), Inches(4.86), Inches(3.35), Inches(0.72), "建议用途：\n体验核心教练对话、职业镜像报告与岗位对比链路。", size=11.5, color=MUTED)
    footer(s, "团队页可继续补充未来核心岗位画像与外部合作顾问信息。")

    prs.save(OUTPUT_PPTX)


if __name__ == "__main__":
    build()
